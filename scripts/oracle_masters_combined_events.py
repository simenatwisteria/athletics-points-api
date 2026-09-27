"""Masters mangekamp: alle klassekolonnene i NFIFs masters-tabeller → fasit.

Skriver ``tests/fixtures/masters_combined_events_cases.json``. Hver klassekolonne (M/W35–M/W100) i
hvert ark blir én blokk med caser resultat → poeng. Seniorkolonnen («Sr») er fasit for WA Combined
Events og ligger i ``wa_combined_events_cases.json`` (AP-012).

Arkene har to oppsett:

- poeng i nøkkelkolonnen og resultat i klassekolonnene (de fleste løp og kast); cellen er det
  resultatet som gir poengene i raden
- resultat i nøkkelkolonnen og poeng i klassekolonnene (hopp og de fleste «-m»-arkene);
  gjenkjennes på «Resultat»/«Res.» i nøkkelkolonnen over dataradene

Hekkearkene har flere tabeller side om side, én per hekkedistanse. Redskap og hekkehøyde per
klasse står i raden «Vekt:»/«Hhøyde» og tas med i hver blokk (BV-034).

Utelatt (BV-024, BV-035, docs/KILDEAVVIK.md), og listet i ``meta``:

- arkene og blokkene for manuell tid på 60 m, 60 m hekk og 80 m hekk, som bruker +0,20 s
- celler som ikke kan leses som tid (menn ``200m`` M95/M100, der et siffer er skjøvet inn i neste
  kolonne, for eksempel «5  1.09.8»)
- avkortede tider i ark med automatisk tid: tekst med bare ett siffer etter siste punktum
  («1.00.0»), der arket ellers har hundredeler. Ingen celler repareres.

Importerer ikke athletics_scoring.

    python scripts/oracle_masters_combined_events.py [--check]

``--check`` skriver ingenting, men feiler hvis fixture-fila i repoet avviker fra det kildene gir.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "tests" / "fixtures" / "masters_combined_events_cases.json"
SOURCES = {
    "M": ROOT / "sources" / "masters" / "nfif-masters-mangekamp-menn.xlsx",
    "F": ROOT / "sources" / "masters" / "nfif-masters-mangekamp-kvinner.xlsm",
}
CLASS_PREFIX = {"M": "M", "F": "W"}
AGES = [str(age) for age in range(35, 101, 5)]

# En blokk: (event_id, nøkkelkolonne, første og siste klassekolonne). Klassene leses fra rad 2.
Block = tuple[str, str, str, str]
STANDARD = ("A", "C", "P")


def _standard(event_id: str) -> list[Block]:
    return [(event_id, *STANDARD)]


# Ark → (blokker, manuell tid, redskapsenhet). Enhet: kg, g (spyd) eller cm (hekkehøyde).
SHEETS: dict[str, dict[str, tuple[list[Block], bool, str | None]]] = {
    "M": {
        "60m": (_standard("sprint_60m"), False, None),
        "100m": (_standard("sprint_100m"), False, None),
        "100m-m": (_standard("sprint_100m"), True, None),
        "200m": (_standard("sprint_200m"), False, None),
        "200m-m": (_standard("sprint_200m"), True, None),
        "400m": (_standard("sprint_400m"), False, None),
        "400m-m": (_standard("sprint_400m"), True, None),
        "1000m": (_standard("middle_1000m"), False, None),
        "1500m": (_standard("middle_1500m"), False, None),
        "60m HK": (_standard("hurdles_60m"), False, "cm"),
        "80-110m HK": (
            [("hurdles_110m", "A", "C", "E"), ("hurdles_100m", "G", "H", "K"),
             ("hurdles_80m", "M", "N", "T")],
            False, "cm",
        ),
        "80-110m HK-m": (
            [("hurdles_110m", "A", "C", "E"), ("hurdles_100m", "G", "H", "K"),
             ("hurdles_80m", "M", "N", "T")],
            True, "cm",
        ),
        "Høyde": (_standard("high_jump"), False, None),
        "Stav": (_standard("pole_vault"), False, None),
        "Lengde": (_standard("long_jump"), False, None),
        "Kule": (_standard("shot_put"), False, "kg"),
        "Diskos": (_standard("discus"), False, "kg"),
        "Spyd": (_standard("javelin"), False, "g"),
        "Slegge": (_standard("hammer"), False, "kg"),
        "Vektkast": (_standard("weight_throw"), False, "kg"),
    },
    "F": {
        "60m": (_standard("sprint_60m"), False, None),
        "100m": (_standard("sprint_100m"), False, None),
        "100m-m": (_standard("sprint_100m"), True, None),
        "200m": (_standard("sprint_200m"), False, None),
        "200m-m": (_standard("sprint_200m"), True, None),
        "400m": (_standard("sprint_400m"), False, None),
        "400m-m": (_standard("sprint_400m"), True, None),
        "800m": (_standard("middle_800m"), False, None),
        "1500m": (_standard("middle_1500m"), False, None),
        "60mHK": (_standard("hurdles_60m"), False, "cm"),
        "80-100mHK": (
            [("hurdles_80m", "A", "B", "M"), ("hurdles_100m", "S", "R", "R")], False, "cm"
        ),
        "80-100mHK-m": (
            [("hurdles_80m", "A", "B", "M"), ("hurdles_100m", "Q", "S", "S")], True, "cm"
        ),
        "Høyde": (_standard("high_jump"), False, None),
        "Lengde": (_standard("long_jump"), False, None),
        "Stav": (_standard("pole_vault"), False, None),
        "Kule": (_standard("shot_put"), False, "kg"),
        "Diskos": (_standard("discus"), False, "kg"),
        "Spyd": (_standard("javelin"), False, "g"),
        "Slegge": (_standard("hammer"), False, "kg"),
        "Vektkast": (_standard("weight_throw"), False, "kg"),
    },
}  # fmt: skip
# Ark som bevisst ikke brukes, med grunn. Alle andre ark enn MKTAB må stå i SHEETS eller her.
SKIPPED: dict[str, dict[str, str]] = {
    "M": {
        "60m-m": "+0,20 s i stedet for +0,24 s (BV-024)",
        "60m HK-m": "+0,20 s i stedet for +0,24 s (BV-024)",
    },
    "F": {
        "60m-m": "+0,20 s i stedet for +0,24 s (BV-024)",
        "60mHK-m": "+0,20 s i stedet for +0,24 s (BV-024)",
    },
}
# Blokker som bevisst ikke brukes: (ark, event_id) → grunn.
SKIPPED_BLOCKS: dict[str, dict[tuple[str, str], str]] = {
    "M": {("80-110m HK-m", "hurdles_80m"): "+0,20 s i stedet for +0,24 s (BV-024)"},
    "F": {("80-100mHK-m", "hurdles_80m"): "+0,20 s i stedet for +0,24 s (BV-024)"},
}
TIME_RE = re.compile(r"^(?:(\d+)\.)?(\d{1,2}\.\d{1,2})$")
TRUNCATED_RE = re.compile(r"^[\d.]*\.\d$")
FIELD_EVENTS = {"high_jump", "pole_vault", "long_jump", "shot_put", "discus", "javelin", "hammer",
                "weight_throw"}  # fmt: skip
IMPLEMENT_LABELS = ("Vekt", "Hhøyde", "HHøyde")


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float):
        return repr(value)
    text = str(value).strip()
    return text or None


def _points(text: str) -> int:
    if not text.isdigit():
        raise ValueError(f"ikke et poengtall: {text!r}")
    return int(text)


def _result(text: str, event_id: str) -> dict[str, float | int]:
    if event_id in FIELD_EVENTS:
        return {"distance_meters": float(Decimal(text))}
    match = TIME_RE.match(text)
    if match is None:
        raise ValueError(f"ikke en tid: {text!r}")
    minutes, seconds = match.groups()
    if minutes is None:
        return {"time_seconds": float(Decimal(seconds))}
    return {"time_minutes": int(minutes), "time_seconds": float(Decimal(seconds))}


def _implement(text: str, unit: str) -> str:
    """'7.26' kg → '7,26kg', '800' g → '0,8kg' (spyd i kg, som Tyrving), '99' cm → '99cm'."""
    value = Decimal(text)
    if unit == "g":
        value, unit = value / 1000, "kg"
    return f"{format(value.normalize(), 'f').replace('.', ',')}{unit}"


def _col(letter: str) -> int:
    return column_index_from_string(letter) - 1


def _unusable(raw: object, text: str, event_id: str, manual: bool) -> str | None:
    """Hvorfor en resultatcelle ikke kan brukes som fasit, eller ``None`` hvis den kan (BV-035)."""
    if event_id in FIELD_EVENTS:
        return None
    if TIME_RE.match(text) is None:
        return "kan ikke leses som tid"
    if not manual and isinstance(raw, str) and TRUNCATED_RE.match(text):
        return "avkortet: ett siffer etter siste punktum i ark med hundredeler"
    return None


def read_block(
    rows: list[tuple[Any, ...]],
    title: str,
    block: Block,
    unit: str | None,
    manual: bool,
    skipped: list[dict[str, Any]],
) -> dict[str, list[Any]]:
    """Klasse → [kolonne, redskap, caser] for én blokk.

    Case: ``[radnummer, resultat i arket, result, poeng]``. Celler som ikke kan brukes, legges i
    ``skipped`` med rad, kolonne og råtekst.
    """
    event_id, key_letter, first, last = block
    key = _col(key_letter)
    header = {c: _text(rows[1][c]) for c in range(_col(first), _col(last) + 1)}
    # Kolonner uten klasse i rad 2 brukes ikke. Har de data, listes de (kvinner 100m-m, kolonne P).
    for c in [c for c, h in header.items() if h is None]:
        filled = [n for n, r in enumerate(rows[3:], start=4) if _text(r[c]) is not None]
        if filled:
            skipped.append(
                {"sheet": title, "rows": f"{filled[0]}–{filled[-1]}", "count": len(filled),
                 "column": get_column_letter(c + 1), "age_class": None,
                 "reason": "kolonnen har data, men ingen klasse i rad 2"}
            )  # fmt: skip
    columns = [c for c, h in header.items() if h is not None]
    if not all(header[c] in AGES for c in columns):
        raise ValueError(f"{title} {first}–{last}: uventede klasser i rad 2: {header}")
    implements: dict[int, str | None] = {c: None for c in columns}
    if unit is not None:
        label = _text(rows[2][0]) or ""
        if not label.startswith(IMPLEMENT_LABELS):
            raise ValueError(f"{title}: rad 3 er ikke «Vekt:»/«Hhøyde»")
        implements = {c: _implement(str(_text(rows[2][c])), unit) for c in columns}
    result_first = any((_text(r[key]) or "").startswith("Res") for r in rows[:5])

    out: dict[str, list[Any]] = {}
    for c in columns:
        cases = []
        for number, row in enumerate(rows, start=1):
            k, v = _text(row[key]), _text(row[c])
            if number <= 3 or k is None or v is None or not re.match(r"^\d", k):
                continue
            result_text, points_text = (k, v) if result_first else (v, k)
            raw = row[key] if result_first else row[c]
            reason = _unusable(raw, result_text, event_id, manual)
            if reason is not None:
                skipped.append(
                    {"sheet": title, "row": number, "column": get_column_letter(c + 1),
                     "age_class": header[c], "text": result_text, "reason": reason}
                )  # fmt: skip
                continue
            result = _result(result_text, event_id)
            cases.append([number, result_text, result, _points(points_text)])
        if not cases:
            raise ValueError(f"{title} kolonne {get_column_letter(c + 1)}: ingen rader")
        out[str(header[c])] = [get_column_letter(c + 1), implements[c], cases]
    return out


def generate() -> dict[str, Any]:
    blocks: list[dict[str, Any]] = []
    skipped_blocks: list[dict[str, Any]] = []
    skipped_cells: list[dict[str, Any]] = []
    for gender, path in SOURCES.items():
        wb = openpyxl.load_workbook(path, data_only=True)
        titles = {ws.title for ws in wb.worksheets} - {"MKTAB"}
        unknown = titles - set(SHEETS[gender]) - set(SKIPPED[gender])
        if unknown or not set(SHEETS[gender]) <= titles:
            raise ValueError(f"{path.name}: uventede eller manglende ark {sorted(unknown)}")
        for title, (sheet_blocks, manual, unit) in SHEETS[gender].items():
            rows = list(wb[title].iter_rows(values_only=True))
            for block in sheet_blocks:
                reason = SKIPPED_BLOCKS[gender].get((title, block[0]))
                if reason is not None:
                    classes = [_text(rows[1][c]) for c in range(_col(block[2]), _col(block[3]) + 1)]
                    skipped_blocks.append(
                        {"gender": gender, "sheet": title, "event_id": block[0],
                         "columns": f"{block[2]}–{block[3]}",
                         "age_classes": [f"{CLASS_PREFIX[gender]}{a}" for a in classes if a],
                         "reason": reason}
                    )  # fmt: skip
                    continue
                cells: list[dict[str, Any]] = []
                read = read_block(rows, title, block, unit, manual, cells)
                skipped_cells += [
                    {"gender": gender, **cell,
                     "age_class": cell["age_class"]
                     and f"{CLASS_PREFIX[gender]}{cell['age_class']}"}
                    for cell in cells
                ]  # fmt: skip
                for age, (column, implement, cases) in read.items():
                    blocks.append(
                        {
                            "file": str(path.relative_to(ROOT)),
                            "sheet": title,
                            "column": column,
                            "gender": gender,
                            "age_class": f"{CLASS_PREFIX[gender]}{age}",
                            "event_id": block[0],
                            "implement": implement,
                            "manual_timing": manual,
                            "cases": cases,
                        }
                    )
    return {
        "meta": {
            "scoring_system": "masters_combined_events",
            "version": "2023",
            "generated_by": "scripts/oracle_masters_combined_events.py",
            "description": "Klassekolonnene (M/W35–100) i NFIFs masters-mangekamptabeller",
            "case_format": "[radnummer, resultat i arket, result (models.Result), poeng]",
            "sources": {
                str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in SOURCES.values()
            },
            "skipped_sheets": {g: dict(s) for g, s in SKIPPED.items()},
            "skipped_blocks": skipped_blocks,
            "skipped_cell_count": len(skipped_cells),
            "skipped_cells": skipped_cells,
            "block_count": len(blocks),
            "case_count": sum(len(b["cases"]) for b in blocks),
            "locked": False,
        },
        "blocks": blocks,
    }


def render(data: dict[str, Any]) -> str:
    """Meta og blokkhoder med innrykk, én case per linje (holder fila liten og diffbar)."""

    def compact(value: Any) -> str:
        return json.dumps(value, ensure_ascii=False, separators=(", ", ": "))

    cells = data["meta"]["skipped_cells"]
    meta = {**data["meta"], "skipped_cells": "@CELLS@"}
    text = json.dumps(meta, ensure_ascii=False, indent=2).replace("\n", "\n  ")
    rows = ",\n".join(f"      {compact(cell)}" for cell in cells)
    text = text.replace('"@CELLS@"', f"[\n{rows}\n    ]" if cells else "[]")
    lines = ["{", f'  "meta": {text},', '  "blocks": [']
    for i, block in enumerate(data["blocks"]):
        head = ", ".join(f"{compact(k)}: {compact(v)}" for k, v in block.items() if k != "cases")
        rows = ",\n".join(f"      {compact(case)}" for case in block["cases"])
        end = "," if i < len(data["blocks"]) - 1 else ""
        lines.append(f'    {{{head}, "cases": [\n{rows}\n    ]}}{end}')
    lines += ["  ]", "}"]
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=(__doc__ or "").splitlines()[0])
    parser.add_argument("--check", action="store_true", help="feil hvis fixture-fila er utdatert")
    args = parser.parse_args()

    data = generate()
    output = render(data)
    if json.loads(output) != data:
        raise ValueError("render() gir ikke samme data tilbake")
    target = TARGET.relative_to(ROOT)
    if args.check:
        if not TARGET.exists() or TARGET.read_text(encoding="utf-8") != output:
            print(f"{target} er utdatert — kjør skriptet uten --check", file=sys.stderr)
            return 1
        print("OK")
        return 0
    TARGET.write_text(output, encoding="utf-8")
    meta = data["meta"]
    print(f"Skrev {target} ({meta['block_count']} blokker, {meta['case_count']} caser)")
    per_sheet: dict[tuple[str, str], int] = {}
    for block in data["blocks"]:
        key = (block["gender"], block["sheet"])
        per_sheet[key] = per_sheet.get(key, 0) + len(block["cases"])
    for (gender, sheet), count in per_sheet.items():
        print(f"  {gender} {sheet}: {count}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
