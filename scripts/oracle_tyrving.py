"""Tyrving-oracle: regnearket regnet om av LibreOffice → tests/fixtures/tyrving_cases.json.

Fasiten kommer fra NFIFs eget regneark, ikke fra athletics_scoring (som dette skriptet aldri
importerer). Framgangsmåte:

1. Kopier regnearket til en temp-mappe (``sources/`` røres aldri).
2. Rett radene der regnearket avviker fra den offisielle PDF-en, i kopien (PATCHES,
   docs/KILDEAVVIK.md).
3. For hver resultatvariant: skriv ett resultat per rad i input-cellene (D/E), og la LibreOffice
   headless regne om hele arbeidsboka (egen profil som tvinger omregning ved lasting).
4. Les poengene i kolonne F. Der mellomresultatet P ligger innenfor 1e-6 fra et heltall, men F er
   rundet ned til heltallet under, merkes casen ``float_edge``: regnearket har da en flyttallsfeil,
   mens PDF-regelen («alle desimaler beholdes») gir heltallet. Fasiten (``points``) er fortsatt
   regnearkets verdi. Hva motoren skal gjøre, avgjøres i review (AP-005).

    python scripts/oracle_tyrving.py [--check]

``--check`` skriver ingenting, men feiler hvis fixture-fila avviker fra det oracle gir.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from extract_tyrving_params import event_id, normalize_spec

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "sources" / "tyrving" / "tyrving-2014-redigerbar.xlsx"
TARGET = ROOT / "tests" / "fixtures" / "tyrving_cases.json"
SHEET_RE = re.compile(r"^(Gutter|Jenter) (\d{2}) år$")

THREE_INTERVAL_P = (
    "=IF(E{r}<0.8*H{r},1000-((N{r}-0.8*N{r})*K{r}+(O{r}-(N{r}-N{r}*0.8))*L{r}),1000-O{r}*I{r})"
)
# Rettinger i kopien der PDF-en vinner. «before» sjekkes mot kilden, så skriptet feiler hvis
# regnearket endres.
PATCHES: list[dict[str, Any]] = [
    {
        "sheet": "Gutter 19 år",
        "reason": "2000 m: multiplikator 0,5 i regnearket, 0,45 i PDF-en",
        "cells": {"I16": {"before": 0.5, "after": 0.45}},
    },
    {
        "sheet": "Jenter 17 år",
        "reason": "Kule 3kg: enkel kvotient med I=1.2 i regnearket, tre-intervall i PDF-en",
        "cells": {
            "I36": {"before": 1.2, "after": "=IF(E36<0.8*H36,L36,Q36)"},
            "P36": {"before": "=1000-O36*I36", "after": THREE_INTERVAL_P.format(r=36)},
        },
    },
    {
        "sheet": "Jenter 15 år",
        "reason": "Spyd: 0,4kg / 1000p = 42 i regnearket, 500 g / 1000p = 38 i PDF-en",
        "cells": {
            "C36": {"before": "0,4kg", "after": "0,5kg"},
            "H36": {"before": 42, "after": 38},
        },
    },
]

# Resultat som andel av 1000p-nivået. 0.05 (distanse) og 3.0 (tid) gir 0 poeng.
SIMPLE_FACTORS = {"time": (0.9, 1.0, 1.1, 3.0), "distance": (0.05, 0.9, 1.0, 1.1)}
THREE_INTERVAL_FACTORS = (0.05, 0.7, 0.8, 0.9, 1.0, 1.1)

# LibreOffice regner ikke om xlsx ved lasting som standard. Profilen tvinger det.
RECALC_PROFILE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""  # noqa: E501


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def libreoffice_version() -> str:
    out = subprocess.run(["soffice", "--version"], capture_output=True, text=True, check=True)
    return " ".join(out.stdout.split()[:2])


def _rows(wb: openpyxl.Workbook) -> list[tuple[str, int]]:
    rows = []
    for ws in wb.worksheets:
        if not SHEET_RE.match(ws.title):
            continue
        for row in range(1, ws.max_row + 1):
            if str(ws[f"F{row}"].value or "").startswith("="):
                rows.append((ws.title, row))
    return rows


def _row_kind(ws: Any, row: int) -> tuple[str, str, int]:
    """(målt størrelse, formeltype, skala) lest fra formlene i raden."""
    p_formula = str(ws[f"P{row}"].value)
    measure = "time" if p_formula.startswith("=1000+") else "distance"
    kind = "three_interval" if "0.8*H" in p_formula else "simple_quotient"
    scale = 10 if str(ws[f"M{row}"].value).endswith(")*10") else 100
    return measure, kind, scale


def _result(h1000: float, factor: float, measure: str, scale: int) -> float:
    """Resultat med oppløsningen tabellen bruker: hundredeler/centimeter, eller tideler."""
    value = h1000 * factor
    decimals = 1 if scale == 10 else 2
    rounded = round(value, decimals)
    return 0.01 if measure == "distance" and rounded == 0 else rounded


def _split_time(seconds: float, scale: int) -> tuple[int | None, float]:
    if scale == 100:
        return None, seconds
    minutes = math.floor(seconds / 60)
    return minutes, round(seconds - minutes * 60, 1)


def _patch(wb: openpyxl.Workbook) -> None:
    for patch in PATCHES:
        ws = wb[patch["sheet"]]
        for cell, change in patch["cells"].items():
            if ws[cell].value != change["before"]:
                raise ValueError(
                    f"{patch['sheet']}!{cell} er {ws[cell].value!r}, forventet {change['before']!r}"
                )
            ws[cell].value = change["after"]


def _recalculate(workbook_path: Path, workdir: Path) -> Path:
    profile = workdir / "profile"
    (profile / "user").mkdir(parents=True, exist_ok=True)
    (profile / "user" / "registrymodifications.xcu").write_text(RECALC_PROFILE, encoding="utf-8")
    outdir = workdir / "out"
    subprocess.run(
        ["soffice", f"-env:UserInstallation={profile.as_uri()}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(outdir), str(workbook_path)],
        check=True,
        capture_output=True,
    )  # fmt: skip
    result = outdir / workbook_path.name
    if not result.exists():
        raise RuntimeError(f"LibreOffice skrev ikke {result}")
    return result


def generate() -> dict[str, Any]:
    if shutil.which("soffice") is None:
        raise RuntimeError("Fant ikke soffice (LibreOffice) i PATH")

    with tempfile.TemporaryDirectory() as tmp:
        workdir = Path(tmp)
        base = openpyxl.load_workbook(SOURCE)
        _patch(base)
        rows = _rows(base)

        # Plan: for hver rad, en liste resultater. Variant k setter resultat k i alle rader.
        plan: dict[tuple[str, int], list[dict[str, Any]]] = {}
        for sheet, row in rows:
            ws = base[sheet]
            measure, kind, scale = _row_kind(ws, row)
            three = kind == "three_interval"
            factors = THREE_INTERVAL_FACTORS if three else SIMPLE_FACTORS[measure]
            h1000 = ws[f"H{row}"].value
            results = []
            for factor in factors:
                value = _result(h1000, factor, measure, scale)
                if measure == "time":
                    minutes, seconds = _split_time(value, scale)
                    results.append({"time_minutes": minutes, "time_seconds": seconds})
                else:
                    results.append({"distance_meters": value})
            plan[(sheet, row)] = results

        variants = max(len(r) for r in plan.values())
        points: dict[tuple[str, int, int], Any] = {}
        p_values: dict[tuple[str, int, int], Any] = {}
        for k in range(variants):
            wb = openpyxl.load_workbook(SOURCE)
            _patch(wb)
            for (sheet, row), results in plan.items():
                ws = wb[sheet]
                ws[f"D{row}"].value = None
                ws[f"E{row}"].value = None
                if k >= len(results):
                    continue
                result = results[k]
                if "distance_meters" in result:
                    ws[f"E{row}"].value = result["distance_meters"]
                else:
                    ws[f"D{row}"].value = result["time_minutes"]
                    ws[f"E{row}"].value = result["time_seconds"]
            path = workdir / f"variant-{k}.xlsx"
            wb.save(path)
            calculated = openpyxl.load_workbook(_recalculate(path, workdir), data_only=True)
            for (sheet, row), results in plan.items():
                if k < len(results):
                    points[(sheet, row, k)] = calculated[sheet][f"F{row}"].value
                    p_values[(sheet, row, k)] = calculated[sheet][f"P{row}"].value

        cases = []
        for sheet, row in rows:
            ws = base[sheet]
            match = SHEET_RE.match(sheet)
            assert match
            name = " ".join(str(ws[f"B{row}"].value).split())
            spec = normalize_spec(ws[f"C{row}"].value)
            for k, result in enumerate(plan[(sheet, row)]):
                value = points[(sheet, row, k)]
                if not isinstance(value, int | float) or value != int(value):
                    raise ValueError(f"{sheet}!F{row} variant {k}: uventet poeng {value!r}")
                case: dict[str, Any] = {
                    "event_id": event_id(name, spec),
                    "gender": {"Gutter": "M", "Jenter": "F"}[match[1]],
                    "age": int(match[2]),
                    "implement": None if spec == "Kappgang" else spec,
                    "result": {k2: v for k2, v in result.items() if v is not None},
                    "points": int(value),
                    "source": {"sheet": sheet, "row": row},
                }
                p_value = p_values[(sheet, row, k)]
                nearest = round(p_value)
                if abs(p_value - nearest) < 1e-6 and nearest > 0 and int(value) == nearest - 1:
                    case["float_edge"] = {"p_libreoffice": p_value, "points_if_exact": nearest}
                cases.append(case)

    return {
        "meta": {
            "scoring_system": "tyrving",
            "version": "2014",
            "generated_by": "scripts/oracle_tyrving.py",
            "oracle": f"{libreoffice_version()} headless, omregning av regnearket",
            "source_file": str(SOURCE.relative_to(ROOT)),
            "source_sha256": _sha256(SOURCE),
            "patches": [
                {"sheet": p["sheet"], "reason": p["reason"], "cells": sorted(p["cells"])}
                for p in PATCHES
            ],
            "patch_policy": "PDF vinner ved avvik fra regnearket (docs/KILDEAVVIK.md)",
            "case_count": len(cases),
            "float_edge_count": sum("float_edge" in c for c in cases),
            "result_fields": "som athletics_scoring.models.Result",
        },
        "cases": cases,
    }


def render(data: dict[str, Any]) -> str:
    return json.dumps(data, ensure_ascii=False, indent=1) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=(__doc__ or "").splitlines()[0])
    parser.add_argument("--check", action="store_true", help="feil hvis fixture-fila er utdatert")
    args = parser.parse_args()
    output = render(generate())
    target = TARGET.relative_to(ROOT)
    if args.check:
        if not TARGET.exists() or TARGET.read_text(encoding="utf-8") != output:
            print(f"{target} avviker fra oracle", file=sys.stderr)
            return 1
        print("OK")
        return 0
    TARGET.write_text(output, encoding="utf-8")
    print(f"Skrev {target} ({json.loads(output)['meta']['case_count']} caser)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
