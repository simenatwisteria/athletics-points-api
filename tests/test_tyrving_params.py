"""Lag 2 (B-16): hver parameter i tyrving_parameters_2014.json kan spores til en kildecelle.

Leser regnearket direkte og importerer ikke ekstraksjonsskriptet, så en feil i skriptet ikke kan
skjule seg. Rader der PDF-en har vunnet (``pdf_override``), sammenlignes via ``pdf_override.excel``.
"""

import hashlib
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from openpyxl.workbook import Workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "sources" / "tyrving" / "tyrving-2014-redigerbar.xlsx"
SOURCE_XLS = ROOT / "sources" / "tyrving" / "tyrving-2014.xls"
PARAMS = ROOT / "athletics_scoring" / "data" / "tyrving_parameters_2014.json"

# Rader der regnearket avviker fra PDF-en og PDF-en vinner (docs/KILDEAVVIK.md).
# Endres lista, må Simen se på det.
EXPECTED_OVERRIDES = {("Gutter 19 år", 16), ("Jenter 17 år", 36), ("Jenter 15 år", 36)}
# Cellene der tyrving-2014.xls avviker fra tyrving-2014-redigerbar.xlsx.
EXPECTED_XLS_DIFFERENCES = {
    ("Jenter 15 år", "C36", "0,4kg", "0,5kg"),
    ("Jenter 15 år", "H36", 42, 38),
}


@pytest.fixture(scope="module")
def data() -> dict[str, Any]:
    result: dict[str, Any] = json.loads(PARAMS.read_text(encoding="utf-8"))
    return result


@pytest.fixture(scope="module")
def workbook() -> Workbook:
    return openpyxl.load_workbook(SOURCE, data_only=False)


def _excel(entry: dict[str, Any], field: str) -> Any:
    """Verdien regnearket har for feltet, også der PDF-en har vunnet."""
    return entry.get("pdf_override", {}).get("excel", {}).get(field, entry[field])


def _formula_rows(workbook: Workbook) -> set[tuple[str, int]]:
    rows = set()
    for ws in workbook.worksheets[1:]:
        for row in range(1, ws.max_row + 1):
            value = ws[f"F{row}"].value
            if isinstance(value, str) and value.startswith("="):
                rows.add((ws.title, row))
    return rows


def test_meta_points_at_unchanged_sources(data: dict[str, Any]) -> None:
    for path, sha256 in data["meta"]["sources"].items():
        assert hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == sha256
    assert data["meta"]["entry_count"] == len(data["entries"]) == 560


def test_every_formula_row_extracted_exactly_once(
    data: dict[str, Any], workbook: Workbook
) -> None:
    sources = [(e["source"]["sheet"], e["source"]["row"]) for e in data["entries"]]
    assert len(sources) == len(set(sources))
    assert set(sources) == _formula_rows(workbook)


def test_combinations_are_unique(data: dict[str, Any]) -> None:
    keys = [(e["event_id"], e["gender"], e["age"], e["implement"]) for e in data["entries"]]
    assert len(keys) == len(set(keys))


def test_overrides_are_exactly_the_known_ones(data: dict[str, Any]) -> None:
    overrides = {
        (e["source"]["sheet"], e["source"]["row"]) for e in data["entries"] if "pdf_override" in e
    }
    assert overrides == EXPECTED_OVERRIDES
    assert data["meta"]["override_count"] == len(EXPECTED_OVERRIDES)
    for entry in data["entries"]:
        if "pdf_override" in entry:
            assert entry["pdf_override"]["excel"], "en overstyring må vise hva regnearket sa"


def test_params_equal_excel_cells(data: dict[str, Any], workbook: Workbook) -> None:
    cells = {"h1000": "H", "quotient": "I", "f1": "J", "f2": "K", "f3": "L"}
    mismatches = []
    for entry in data["entries"]:
        ws = workbook[entry["source"]["sheet"]]
        row = entry["source"]["row"]
        for name, value in _excel(entry, "params").items():
            cell = ws[f"{cells[name]}{row}"].value
            if cell != value or isinstance(cell, str):
                mismatches.append((ws.title, row, name, value, cell))
    assert not mismatches


def test_identity_matches_sheet_and_row(data: dict[str, Any], workbook: Workbook) -> None:
    for entry in data["entries"]:
        ws = workbook[entry["source"]["sheet"]]
        row = entry["source"]["row"]
        match = re.fullmatch(r"(Gutter|Jenter) (\d{2}) år", ws.title)
        assert match
        assert entry["gender"] == {"Gutter": "M", "Jenter": "F"}[match[1]]
        assert entry["age"] == int(match[2])
        assert entry["name"] == " ".join(str(ws[f"B{row}"].value).split())
        spec = ws[f"C{row}"].value
        expected = None if spec in (None, "Kappgang") else re.sub(r"\s+", "", str(spec))
        assert _excel(entry, "implement") == expected
        assert entry["event_id"].startswith("racewalk_") == (spec == "Kappgang")


def test_formula_type_matches_excel_formula(data: dict[str, Any], workbook: Workbook) -> None:
    for entry in data["entries"]:
        ws = workbook[entry["source"]["sheet"]]
        p_formula = str(ws[f"P{entry['source']['row']}"].value)
        excel_type = "three_interval" if "0.8*H" in p_formula else "simple_quotient"
        assert _excel(entry, "formula_type") == excel_type


def test_measure_and_scale_match_excel(data: dict[str, Any], workbook: Workbook) -> None:
    for entry in data["entries"]:
        ws = workbook[entry["source"]["sheet"]]
        row = entry["source"]["row"]
        m_formula = str(ws[f"M{row}"].value)
        p_formula = str(ws[f"P{row}"].value)
        assert entry["scale"] == (10 if m_formula.endswith(")*10") else 100)
        assert entry["measure"] == ("time" if p_formula.startswith("=1000+") else "distance")


def test_params_match_formula_type(data: dict[str, Any]) -> None:
    for entry in data["entries"]:
        expected = (
            {"h1000", "f1", "f2", "f3"}
            if entry["formula_type"] == "three_interval"
            else {"h1000", "quotient"}
        )
        assert set(entry["params"]) == expected


def test_json_is_up_to_date() -> None:
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "extract_tyrving_params.py"), "--check"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(shutil.which("soffice") is None, reason="krever LibreOffice (soffice)")
def test_xls_has_same_parameters(workbook: Workbook, tmp_path: Path) -> None:
    """Kryssjekk: den eldre .xls-varianten har samme parametre og formler som .xlsx."""
    shutil.copy(SOURCE_XLS, tmp_path / SOURCE_XLS.name)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path),
         str(tmp_path / SOURCE_XLS.name)],
        check=True,
        capture_output=True,
    )  # fmt: skip
    xls = openpyxl.load_workbook(tmp_path / "tyrving-2014.xlsx", data_only=False)
    assert xls.sheetnames == workbook.sheetnames
    differences = set()
    for sheet, row in sorted(_formula_rows(workbook)):
        for col in "BCHIJKLP":
            a, b = workbook[sheet][f"{col}{row}"].value, xls[sheet][f"{col}{row}"].value
            if a != b:
                differences.add((sheet, f"{col}{row}", a, b))
    assert differences == EXPECTED_XLS_DIFFERENCES
